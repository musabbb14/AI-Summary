"""Tests for DOCX extraction using a document built with python-docx itself."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from docx import Document
from utils.docx_extractor import extract_text_from_docx


def test_extract_docx_paragraphs():
    doc = Document()
    doc.add_paragraph("First paragraph of the test document.")
    doc.add_paragraph("Second paragraph with more content.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    doc.save(path)

    try:
        text = extract_text_from_docx(path)
        assert "First paragraph" in text
        assert "Second paragraph" in text
    finally:
        os.remove(path)
